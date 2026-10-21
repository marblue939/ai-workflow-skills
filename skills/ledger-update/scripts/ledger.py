#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""台账安全更新：先备份、按公司名动态定位、写后回读。

用法：
  python ledger.py backup
  python ledger.py rows [关键词]
  python ledger.py set <公司关键词> <列字母> <新文本>
"""
import argparse, datetime, os, shutil, sys

try:
    import openpyxl
except ImportError:
    sys.exit("需要 openpyxl: pip install openpyxl")

SHEET = "Offer对比总表"          # 按自己的表改
FIRST_DATA_ROW = 5              # 表头之后第一行
NAME_COL = 2                    # B 列 = 公司名
STATUS_COL = 14                 # N 列 = 状态

def data_last_row(ws):
    for r in range(ws.max_row, FIRST_DATA_ROW, -1):
        if str(ws.cell(r, NAME_COL).value or "").startswith("用法"):
            return r - 1
    return ws.max_row

def find_rows(ws, kw):
    return [r for r in range(FIRST_DATA_ROW, data_last_row(ws) + 1)
            if kw in str(ws.cell(r, NAME_COL).value or "")]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["backup", "rows", "set"])
    ap.add_argument("args", nargs="*")
    ap.add_argument("--file", default="台账.xlsx")
    a = ap.parse_args()

    path = os.path.abspath(a.file)
    if a.cmd == "backup":
        stem = os.path.splitext(path)[0] + "_backup_" + datetime.date.today().isoformat()
        cand, i = stem + ".xlsx", ord("b")
        while os.path.exists(cand):
            cand = stem + chr(i) + ".xlsx"; i += 1
        shutil.copy2(path, cand)
        print("BACKUP ->", os.path.basename(cand)); return

    wb = openpyxl.load_workbook(path); ws = wb[SHEET]

    if a.cmd == "rows":
        kw = a.args[0] if a.args else ""
        for r in range(FIRST_DATA_ROW, data_last_row(ws) + 1):
            if kw and kw not in str(ws.cell(r, NAME_COL).value or ""):
                continue
            print(r, "|", ws.cell(r, NAME_COL).value, "|", ws.cell(r, STATUS_COL).value)
        return

    if len(a.args) != 3:
        sys.exit("用法: set <公司关键词> <列字母> <新文本>")
    kw, col, text = a.args
    hits = find_rows(ws, kw)
    if len(hits) != 1:
        sys.exit(f"命中 {len(hits)} 行，中止：{[ws.cell(r, NAME_COL).value for r in hits]}")
    r = hits[0]
    ci = openpyxl.utils.column_index_from_string(col)
    old = ws.cell(r, ci).value
    shutil.copy2(path, path.replace(".xlsx", "_backup_" + datetime.date.today().isoformat() + ".xlsx"))
    ws.cell(r, ci).value = text
    wb.save(path)
    wb2 = openpyxl.load_workbook(path)
    print(f"OK row={r} 公司={ws.cell(r, NAME_COL).value}")
    print("OLD:", old)
    print("NEW:", wb2[SHEET].cell(r, ci).value)

if __name__ == "__main__":
    main()
